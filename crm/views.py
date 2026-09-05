from django.contrib.auth.models import User
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin

from django.http import HttpResponse
from django.views import View

from django.shortcuts import get_object_or_404
from django.db.models import Count  # Para estadísticas
from django.db.models.functions import TruncMonth  # agrupa fechas por mes

from django.db.models import Q  # Q permite armar condiciones OR/AND complejas
from django.urls import reverse_lazy

import openpyxl
from openpyxl.styles import Font

from django.views.generic import (
    ListView,
    DetailView,
    CreateView,
    UpdateView,
    DeleteView,
    TemplateView,
)

from .models import Client, Interaction, Company

from .forms import (
    ClientForm,
    InteractionForm,
    ComercialForm,
    ComercialUpdateForm,
    CompanyForm,
)


class LoginRequiredView(LoginRequiredMixin):
    # Cualquier vista que herede de esta exige login
    login_url = "login"
    redirect_field_name = "next"


class AdminRequiredView(LoginRequiredView, UserPassesTestMixin):
    # Hereda login obligatorio y agrega una verificación extra (test_func)
    def test_func(self):
        return self.request.user.is_staff


# ---- Vistas de Client ----


class ClientListView(LoginRequiredView, ListView):
    model = Client
    template_name = "crm/client_list.html"
    context_object_name = "clients"

    def get_queryset(self):
        # Filtrar según la búsqueda
        query = self.request.GET.get("q", "")
        if query:
            return Client.objects.filter(
                Q(nombre__icontains=query)
                | Q(apellido__icontains=query)
                | Q(email__icontains=query)
                | Q(company__nombre__icontains=query)
            )
        return Client.objects.all()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["query"] = self.request.GET.get("q", "")
        return context


class ClientDetailView(LoginRequiredView, DetailView):
    model = Client
    template_name = "crm/client_detail.html"
    context_object_name = "client"


class ClientCreateView(LoginRequiredView, CreateView):
    model = Client
    form_class = ClientForm
    template_name = "crm/client_form.html"
    success_url = reverse_lazy("client_list")


class ClientUpdateView(LoginRequiredView, UpdateView):
    model = Client
    form_class = ClientForm
    template_name = "crm/client_form.html"
    success_url = reverse_lazy("client_list")


class ClientDeleteView(LoginRequiredView, DeleteView):
    model = Client
    template_name = "crm/client_confirm_delete.html"
    success_url = reverse_lazy("client_list")


from django.http import HttpResponse
from django.views import View
import openpyxl
from openpyxl.styles import Font


class ClientExportView(LoginRequiredView, View):
    def get(self, request):
        # reutilizamos la misma lógica de filtro que ClientListView,
        # así el Excel exportado coincide con lo que el usuario está viendo
        query = request.GET.get("q", "")
        if query:
            clients = Client.objects.filter(
                Q(nombre__icontains=query)
                | Q(apellido__icontains=query)
                | Q(email__icontains=query)
                | Q(company__nombre__icontains=query)
            )
        else:
            clients = Client.objects.all()

        # Workbook: representa el archivo Excel completo en memoria
        wb = openpyxl.Workbook()
        # active: la primera hoja del archivo, la única que necesitamos acá
        ws = wb.active
        ws.title = "Clientes"

        # fila de encabezados, en negrita para diferenciarla de los datos
        headers = [
            "Nombre",
            "Apellido",
            "Email",
            "Teléfono",
            "Empresa",
            "Comercial",
            "Fecha de alta",
        ]
        ws.append(headers)
        for cell in ws[
            1
        ]:  # ws[1] es la primera fila (los encabezados que acabamos de agregar)
            cell.font = Font(bold=True)

        # una fila de datos por cada cliente
        for client in clients:
            ws.append(
                [
                    client.nombre,
                    client.apellido,
                    client.email,
                    client.telefono,
                    str(client.company) if client.company else "",
                    str(client.comercial) if client.comercial else "",
                    # strftime formatea la fecha sin la hora, más limpio para una planilla
                    client.fecha_alta.strftime("%d/%m/%Y"),
                ]
            )

        # ajustamos el ancho de cada columna automáticamente según su contenido más largo,
        # para que no queden columnas angostas con texto cortado
        for col_cells in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col_cells)
            col_letter = col_cells[0].column_letter
            ws.column_dimensions[col_letter].width = max_length + 2

        # HttpResponse con el content_type específico de Excel le indica al navegador
        # que descargue el archivo en vez de intentar mostrarlo como texto plano
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        # Content-Disposition "attachment" fuerza la descarga con el nombre de archivo indicado
        response["Content-Disposition"] = 'attachment; filename="clientes.xlsx"'
        # wb.save() normalmente escribe a un archivo en disco, pero también acepta
        # cualquier objeto tipo "archivo" — HttpResponse funciona como uno
        wb.save(response)
        return response


class InteractionExportView(LoginRequiredView, View):
    # Exportar a Excel,
    def get(self, request):
        interactions = Interaction.objects.select_related("client", "comercial").all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Interacciones"

        headers = ["Cliente", "Tipo", "Comercial", "Fecha", "Notas"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for interaction in interactions:
            ws.append(
                [
                    str(interaction.client),
                    interaction.get_tipo_display(),
                    str(interaction.comercial) if interaction.comercial else "",
                    interaction.fecha.strftime("%d/%m/%Y %H:%M"),
                    interaction.notas,
                ]
            )

        for col_cells in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col_cells)
            col_letter = col_cells[0].column_letter
            ws.column_dimensions[col_letter].width = max_length + 2

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="interacciones.xlsx"'
        wb.save(response)
        return response


# ---- Vistas de Empresa ----


class CompanyListView(LoginRequiredView, ListView):
    model = Company
    template_name = "crm/company_list.html"
    context_object_name = "companies"

    def get_queryset(self):
        query = self.request.GET.get("q", "")
        if query:
            return Company.objects.filter(
                Q(nombre__icontains=query) | Q(sector__icontains=query)
            )
        return Company.objects.all()

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["query"] = self.request.GET.get("q", "")
        return context


class CompanyDetailView(LoginRequiredView, DetailView):
    model = Company
    template_name = "crm/company_detail.html"
    context_object_name = "company"


class CompanyCreateView(LoginRequiredView, CreateView):
    model = Company
    form_class = CompanyForm
    template_name = "crm/company_form.html"
    success_url = reverse_lazy("company_list")


class CompanyUpdateView(LoginRequiredView, UpdateView):
    model = Company
    form_class = CompanyForm
    template_name = "crm/company_form.html"
    success_url = reverse_lazy("company_list")


class CompanyDeleteView(LoginRequiredView, DeleteView):
    model = Company
    template_name = "crm/company_confirm_delete.html"
    success_url = reverse_lazy("company_list")


# ---- Vistas de Interaction ----


class InteractionCreateView(LoginRequiredView, CreateView):
    model = Interaction
    form_class = InteractionForm
    template_name = "crm/interaction_form.html"

    def form_valid(self, form):
        client = get_object_or_404(Client, pk=self.kwargs["client_pk"])
        form.instance.client = client
        self.client_obj = client
        return super().form_valid(form)

    def get_success_url(self):
        return reverse_lazy("client_detail", kwargs={"pk": self.client_obj.pk})

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["client"] = get_object_or_404(Client, pk=self.kwargs["client_pk"])
        return context


class InteractionDeleteView(LoginRequiredView, DeleteView):
    model = Interaction
    template_name = "crm/interaction_confirm_delete.html"

    def get_success_url(self):
        return reverse_lazy("client_detail", kwargs={"pk": self.object.client.pk})


# ---- Vistas de Comercial ----


class ComercialListView(AdminRequiredView, ListView):
    model = User
    template_name = "crm/comercial_list.html"
    context_object_name = "comerciales"

    def get_queryset(self):
        query = self.request.GET.get("q", "")
        base = User.objects.all().order_by("username")
        if query:
            return base.filter(
                Q(username__icontains=query)
                | Q(email__icontains=query)
                | Q(first_name__icontains=query)
                | Q(last_name__icontains=query)
            )
        return base

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["query"] = self.request.GET.get("q", "")
        return context


class ComercialCreateView(AdminRequiredView, CreateView):
    model = User
    form_class = ComercialForm
    template_name = "crm/comercial_form.html"
    success_url = reverse_lazy("comercial_list")


class ComercialUpdateView(AdminRequiredView, UpdateView):
    model = User
    form_class = ComercialUpdateForm
    template_name = "crm/comercial_form.html"
    success_url = reverse_lazy("comercial_list")


class ComercialDeactivateView(AdminRequiredView, UpdateView):
    model = User
    fields = []  # No muestra campos editables. Solo pide confirmación
    template_name = "crm/comercial_confirm_deactivate.html"
    success_url = reverse_lazy("comercial_list")

    def form_valid(self, form):
        form.instance.is_active = False
        return super().form_valid(form)


class MenuView(LoginRequiredView, TemplateView):
    template_name = "crm/menu.html"


# ---- Vistas de Estadísticas ----


class StatsView(AdminRequiredView, TemplateView):
    template_name = "crm/stats.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        import json

        # ---- Interacciones por comercial ----
        comerciales_stats = (
            User.objects.annotate(total_interacciones=Count("interacciones_realizadas"))
            .filter(total_interacciones__gt=0)
            .order_by("-total_interacciones")
        )
        context["comerciales_stats"] = comerciales_stats
        context["labels_comerciales_json"] = json.dumps(
            [c.username for c in comerciales_stats]
        )
        context["data_comerciales_json"] = json.dumps(
            [c.total_interacciones for c in comerciales_stats]
        )

        # ---- Interacciones por cliente ----
        clientes_stats = (
            Client.objects.annotate(total_interacciones=Count("interactions"))
            .filter(total_interacciones__gt=0)
            .order_by("-total_interacciones")
        )
        context["clientes_stats"] = clientes_stats
        context["labels_clientes_json"] = json.dumps(
            [f"{c.nombre} {c.apellido}" for c in clientes_stats]
        )
        context["data_clientes_json"] = json.dumps(
            [c.total_interacciones for c in clientes_stats]
        )

        # ---- Interacciones por mes ----
        interacciones_por_mes = (
            Interaction.objects.annotate(mes=TruncMonth("fecha"))
            .values("mes")  # agrupamos por mes
            .annotate(
                total=Count("id")
            )  # contamos cuántas interacciones caen en cada mes
            .order_by("mes")  # orden cronológico, de más antiguo a más reciente
        )
        context["interacciones_por_mes"] = interacciones_por_mes
        context["labels_meses_json"] = json.dumps(
            [item["mes"].strftime("%b %Y") for item in interacciones_por_mes]
        )
        context["data_meses_json"] = json.dumps(
            [item["total"] for item in interacciones_por_mes]
        )

        return context
